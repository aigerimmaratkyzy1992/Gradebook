import io
from typing import Annotated
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from app.api.deps import Db, require_role
from app.models import Course, User
from app.roles import ADMIN, TEACHER
from app.schemas.gradebook import (
    CourseStatsResponse,
    GradebookCell,
    GradebookColumn,
    GradebookResponse,
    GradebookRow,
    HardestAssignment,
    RankingEntry,
)
from app.schemas.user import UserResponse
from app.services import course as course_service
from app.services import gradebook as gradebook_service

router = APIRouter()

StaffUser = Annotated[User, Depends(require_role(ADMIN, TEACHER))]

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


async def _course(db: Db, course_id: int, user: User) -> Course:
    course = await course_service.get_course(db, course_id)
    if course is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Дисциплина не найдена")
    if user.role != ADMIN and course.teacher_id != user.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Недостаточно прав")
    return course


@router.get("/{course_id}/gradebook")
async def get_gradebook(
    course_id: int, db: Db, user: StaffUser
) -> GradebookResponse:
    course = await _course(db, course_id, user)
    book = await gradebook_service.build(db, course)

    return GradebookResponse(
        columns=[GradebookColumn(**column.__dict__) for column in book.columns],
        rows=[
            GradebookRow(
                student=UserResponse.model_validate(row.student),
                cells=[GradebookCell(**cell.__dict__) for cell in row.cells],
                average=row.average,
                total=row.total,
            )
            for row in book.rows
        ],
    )


@router.get("/{course_id}/stats")
async def get_stats(course_id: int, db: Db, user: StaffUser) -> CourseStatsResponse:
    course = await _course(db, course_id, user)
    book = await gradebook_service.build(db, course)

    students = len(book.rows)
    assignments = len(book.columns)
    expected = students * assignments

    submitted = sum(
        1 for row in book.rows for cell in row.cells if cell.status is not None
    )
    reviewed = sum(
        1
        for row in book.rows
        for cell in row.cells
        if cell.status in ("reviewed", "revision")
    )
    all_scores = [
        cell.score for row in book.rows for cell in row.cells if cell.score is not None
    ]

    # Ең қиын тапсырмалар — орташа пайызы төмендері.
    hardest: list[HardestAssignment] = []
    for index, column in enumerate(book.columns):
        scores = [
            row.cells[index].score
            for row in book.rows
            if row.cells[index].score is not None
        ]
        if scores:
            percent = sum(scores) / len(scores) / column.max_score * 100  # type: ignore[arg-type]
            hardest.append(
                HardestAssignment(
                    id=column.id,
                    title=column.title,
                    average_percent=round(percent, 1),
                )
            )
    hardest.sort(key=lambda item: item.average_percent)

    ranking = [
        RankingEntry(
            student=UserResponse.model_validate(row.student), average=row.average
        )
        for row in book.rows
        if row.average is not None
    ]
    ranking.sort(key=lambda item: item.average, reverse=True)

    return CourseStatsResponse(
        students=students,
        assignments=assignments,
        expected=expected,
        submitted=submitted,
        reviewed=reviewed,
        not_submitted=max(0, expected - submitted),
        average=round(sum(all_scores) / len(all_scores), 1) if all_scores else None,
        completion_percent=round(submitted / expected * 100, 1) if expected else 0.0,
        hardest=hardest[:5],
        ranking=ranking,
    )


@router.get("/{course_id}/gradebook.xlsx")
async def export_gradebook(
    course_id: int, db: Db, user: StaffUser
) -> StreamingResponse:
    course = await _course(db, course_id, user)
    book = await gradebook_service.build(db, course)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал"

    header = ["Студент", *(column.title for column in book.columns), "Средний балл"]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in book.rows:
        sheet.append(
            [
                row.student.full_name or row.student.email,
                *(cell.score if cell.score is not None else "—" for cell in row.cells),
                row.average if row.average is not None else "—",
            ]
        )

    sheet.column_dimensions["A"].width = 32
    for index in range(len(book.columns) + 1):
        sheet.column_dimensions[chr(ord("B") + index)].width = 16

    # Ведомость: әр студент бойынша жинақ.
    summary = workbook.create_sheet("Ведомость")
    summary.append(
        ["Студент", "Сдано работ", "Сумма баллов", "Средний балл", "Успеваемость"]
    )
    total_max = sum(column.max_score for column in book.columns)
    for row in book.rows:
        done = sum(1 for cell in row.cells if cell.status is not None)
        summary.append(
            [
                row.student.full_name or row.student.email,
                f"{done} из {len(book.columns)}",
                row.total,
                row.average if row.average is not None else "—",
                f"{round(row.total / total_max * 100, 1)}%" if total_max else "—",
            ]
        )

    # Список студентов — топтар бойынша.
    people = workbook.create_sheet("Студенты")
    people.append(["№", "ФИО", "Email", "Группа"])
    number = 0
    for group in course.groups:
        for student in group.students:
            number += 1
            people.append(
                [number, student.full_name or "—", student.email, group.name]
            )

    for extra, widths in ((summary, (32, 16, 16, 16, 16)), (people, (6, 32, 28, 16))):
        for cell in extra[1]:
            cell.font = Font(bold=True)
        for letter, width in zip("ABCDE", widths, strict=False):
            extra.column_dimensions[letter].width = width

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    filename = quote(f"Журнал — {course.name}.xlsx")
    return StreamingResponse(
        stream,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{filename}"
        },
    )
